#!/usr/bin/env python3
"""Build the v58 manuscript by editing the IEEE v57 .docx in place.

The v57 build scripts no longer exist; the v57 .docx does, and it is the
better starting point anyway: everything already verified in it (masthead,
running heads, footers, styles, the table grids fixed at v49, biographies)
is left untouched, and only the paragraphs, tables, figure and references
that the paired campaign changes are edited.

    python3 docs/build_v58.py --v57 MultiModelDatabase_AccessControl_IEEE_v57.docx \\
        --p30 analysis/paired_p0.30.json --p10 analysis/paired_p0.10.json \\
        --p05 analysis/paired_p0.05.json --fig6 docs/figures/fig7_ablation.png \\
        --out MultiModelDatabase_AccessControl_IEEE_v58.docx

    python3 docs/build_v58.py --v57 ... --roundtrip     # apply nothing; output text must equal input

Discipline, the same as v57's: every edit names its paragraph by exact text
and must match exactly one paragraph or the build aborts; every number comes
from the ablation_table.py JSON through docs/fill_revision.values_for; a
placeholder left unresolved aborts; and the finished document is checked for
stale phrases, citation order, table and figure sequence, and the abstract's
word limit before it is written.
"""
import argparse
import copy
import re
import struct
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from docs import fill_revision as fr  # noqa: E402
from experiments.analyze import wilson_interval  # noqa: E402
from docs.v58_edits import EDITS, REF_REPLACE, REF_DELETE, REF_INSERT_AFTER, TABLE_VII, TABLE_VIII  # noqa: E402

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]


def die(msg):
    raise SystemExit(f"BUILD ABORTED: {msg}")


# ---------------------------------------------------------------------------
# Values and placeholder resolution
# ---------------------------------------------------------------------------

def load_values(p30, p10, p05):
    v = {}
    reports = {}
    for tag, path in (("p30", p30), ("p10", p10), ("p05", p05)):
        if not path:
            continue
        rep, block, _ = fr.load_block(path)
        reports[tag] = (rep, block)
        v.update(fr.values_for(tag, block, rep))
    b30 = reports["p30"][1]["configs"]
    for tag, (rep, blk) in reports.items():
        for short, label in fr.CFG.items():
            c = blk["configs"].get(label)
            if not c:
                continue
            n, k = c["n"], c["leaked_without_bound"]
            lo, hi = wilson_interval(n - k, n)
            v[f"{tag}.{short}.closed_rate"] = f"{(n - k) / n:.2f}"
            # spelled out below ten, as IEEE prose has it ("six trials")
            words = ["zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine"]
            v[f"{tag}.{short}.k_word"] = words[k] if k < 10 else str(k)
            v[f"{tag}.{short}.closed_lo"] = f"{lo:.2f}"
            v[f"{tag}.{short}.closed_hi"] = f"{hi:.2f}"
            if c["drift_window_median_s"] is not None:
                v[f"{tag}.{short}.win_ms"] = f"{c['drift_window_median_s'] * 1000:.0f}"
                v[f"{tag}.{short}.win_2"] = f"{c['drift_window_median_s']:.2f}"
    v["ledger_extra_s"] = f"{b30['C ledger']['revoke_latency_median_s'] - b30['C log']['revoke_latency_median_s']:.1f}"
    v["log_extra_s"] = f"{b30['C log']['revoke_latency_median_s'] - b30['C retry']['revoke_latency_median_s']:.2f}"
    cr_leaked = b30["C retry"]["leaked_seeds"]
    v["p30.Cr.leaked_seeds"] = fr.oxford(cr_leaked) if cr_leaked else "none"
    o_win = {t["seed"]: t["drift_window_s"] for t in b30["O"]["trials"]}
    wins = [o_win.get(sd) for sd in cr_leaked]
    if cr_leaked and all(w is not None for w in wins):
        v["p30.O.win_leaked_seeds"] = ("between " + f"{min(wins):.2f} and {max(wins):.2f}") if len(wins) > 2 \
            else fr.oxford(f"{w:.2f}" for w in wins)
    elif cr_leaked:
        v["p30.O.win_leaked_seeds"] = "[[O DID NOT CLOSE ON EVERY SUCH SEED -- WRONG BRANCH]]"
    else:
        v["p30.O.win_leaked_seeds"] = "(no such seed)"
    # Injected seeds per p from the schedule itself (not from Scenario A's
    # valid trials, which lose a seed to any exclusion).
    from scenarios.common import faults
    from experiments.run_experiment import fault_seed
    for tag, (rep, blk) in reports.items():
        pv = float(blk["configs"]["A"]["fault_p"][0])
        seeds = sorted({t["seed"] for c in blk["configs"].values() for t in c["trials"]})
        per_layer = {"relational": 0, "nosql": 0, "vector": 0}
        n_inj = 0
        for sd in seeds:
            fs = fault_seed(10000, sd)
            hit = [l for l in per_layer if faults.paired_fault_u(fs, l, 0) < pv]
            if hit:
                n_inj += 1
                for l in hit:
                    per_layer[l] += 1
        v[f"{tag}.injected_n"] = str(n_inj)
        v[f"{tag}.injected_layers"] = ", ".join(f"{l} in {c}" for l, c in sorted(per_layer.items(), key=lambda x: -x[1]))
        # Scenario A must have leaked on every valid injected trial
        a_leaked = blk["configs"]["A"]["leaked_without_bound"]
        a_valid_injected = sum(1 for t in blk["configs"]["A"]["trials"]
                               if any(faults.paired_fault_u(fault_seed(10000, t["seed"]), l, 0) < pv for l in per_layer))
        if a_leaked != a_valid_injected:
            die(f"{tag}: A leaked on {a_leaked} trials but {a_valid_injected} valid trials had an injection")
    # host stalls / exclusions across the three p
    ex = []
    for tag, (rep, blk) in reports.items():
        for label, seeds in blk.get("excluded_seeds", {}).items():
            for seed, reason in seeds:
                ex.append((tag, label, seed, reason))
    n_ex = len(ex)
    unknown = [e for e in ex if e[3] != "poller_error"]
    if unknown:
        # the wording below describes a host stall caught by the clock-anomaly
        # guard; any other exclusion needs its own sentence, written by a person
        die(f"exclusions other than the clock-anomaly guard need authoring: {unknown}")
    if n_ex:
        words = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five"}
        v["host_stall_lead"] = (f"The host stalled during the paired campaign, and {words.get(n_ex, n_ex)} trial"
                                f"{'s were' if n_ex != 1 else ' was'} excluded for it. ")
        v["host_stall_tail"] = "The exclusions are listed in Table XI."
        v["host_stall_each"] = "The affected trial" if n_ex == 1 else "Each affected trial"
        # ablation_table.py records the guard that fired as "poller_error";
        # the poller's own message for the stall is a clock anomaly
        # (verify_v58.py checks that message against this label).
        reason_label = {"poller_error": "clock anomaly"}
        v["excluded_seeds_caption"] = "Excluded: " + "; ".join(
            f"{label} seed {seed} at p = {tag[1:].lstrip('0') and '0.' + tag[1:]} "
            f"({reason_label.get(reason, reason.replace('_', ' '))})"
            for tag, label, seed, reason in ex) + "."
    else:
        v["host_stall_lead"] = ("The host did not stall during the paired campaign; it had stalled twice during "
                                "the campaign behind the first version of this study, for about 78 and 127 seconds. ")
        v["host_stall_tail"] = "No trial in Table VII or Table VIII was excluded for this reason."
        v["host_stall_each"] = "Each affected trial"
        v["excluded_seeds_caption"] = "No trial was excluded."
    return v, reports


_IF = re.compile(r"\{\{IF ([^}]+)\}\}(.*?)(?:\{\{ELSE\}\}(.*?))?\{\{ENDIF\}\}", re.S)


def resolve(text, v, chosen):
    def cond(expr):
        m = re.match(r"([\w.\-]+) (==|>) (\d+)$", expr.strip())
        if not m:
            die(f"cannot parse condition {expr!r}")
        key, op, rhs = m.groups()
        if key not in v:
            die(f"condition refers to {key}, not in the reports")
        try:
            lhs = float(v[key])
        except ValueError:
            die(f"condition on {key} = {v[key]!r}, not a number")
        return lhs == float(rhs) if op == "==" else lhs > float(rhs)

    def branch(m):
        c = cond(m.group(1))
        chosen.append(f"IF {m.group(1)} -> {'IF' if c else 'ELSE'}")
        return (m.group(2) if c else (m.group(3) or ""))
    text = _IF.sub(branch, text)
    unresolved = []

    def sub(m):
        k = m.group(1)
        if k in v:
            return v[k]
        unresolved.append(k)
        return m.group(0)
    text = re.sub(r"\{\{([^}]+)\}\}", sub, text)
    if unresolved:
        die(f"unresolved placeholders: {sorted(set(unresolved))}")
    return text


# ---------------------------------------------------------------------------
# Paragraph helpers
# ---------------------------------------------------------------------------

def find_one(paras, pred, what):
    hits = [p for p in paras if pred(p)]
    if len(hits) != 1:
        die(f"{what}: expected exactly one matching paragraph, found {len(hits)}")
    return hits[0]


def set_run_text(r, text):
    """Assign only when the text changes, and never on a run that carries a
    drawing: python-docx's run.text setter removes every child of the run,
    drawings included (the IEEE v27 masthead loss)."""
    if text == r.text:
        return
    if list(r._r.iter(qn("w:drawing"))):
        die(f"refusing to rewrite a run that holds a drawing: {r.text[:40]!r}")
    r.text = text


def set_single_run(p, text):
    runs = p.runs
    if len(runs) != 1:
        die(f"paragraph is not single-run ({len(runs)} runs): {p.text[:60]!r}")
    set_run_text(runs[0], text)


def remove_paragraph(p):
    p._p.getparent().remove(p._p)


# ---------------------------------------------------------------------------
# Edits
# ---------------------------------------------------------------------------

def apply_edits(doc, v, chosen, log):
    for e in EDITS:
        kind = e["kind"]
        paras = doc.paragraphs
        if kind == "abstract":
            p = find_one(paras, lambda p: p.style.name == "Abstract", "abstract")
            body = p.runs[1] if len(p.runs) == 2 and p.runs[0].text.strip() == "ABSTRACT" else None
            if body is None:
                die("abstract paragraph is not [label, body]")
            if body.text.count(e["find"]) != 1:
                die(f"abstract: 'find' text not found exactly once: {e['find'][:60]!r}")
            set_run_text(body, body.text.replace(e["find"], resolve(e["new"], v, chosen)))
            log.append("abstract: replaced one sentence")
        elif kind == "para":
            if "old" in e:
                p = find_one(paras, lambda p, o=e["old"]: p.text == o, f"para {e['old'][:50]!r}")
            else:
                p = find_one(paras, lambda p, o=e["old_prefix"]: p.text.startswith(o), f"para {e['old_prefix'][:50]!r}")
            set_single_run(p, resolve(e["new"], v, chosen))
            log.append(f"para: {p.text[:50]!r}")
        elif kind == "para_sub":
            p = find_one(paras, lambda p, o=e["old_prefix"]: p.text.startswith(o), f"para_sub {e['old_prefix'][:50]!r}")
            if len(p.runs) != 1:
                die(f"para_sub target is not single-run: {p.text[:50]!r}")
            if p.runs[0].text.count(e["find"]) != 1:
                die(f"para_sub: 'find' not found exactly once in {p.text[:50]!r}: {e['find'][:60]!r}")
            set_run_text(p.runs[0], p.runs[0].text.replace(e["find"], resolve(e["new"], v, chosen)))
            log.append(f"para_sub: {e['find'][:40]!r}")
        elif kind in ("bullet", "bullet_sub"):
            def is_target(p, e=e):
                if len(p.runs) != 3 or p.runs[0].text != "• " or p.runs[1].text != e["old_lead"]:
                    return False
                if "old" in e:
                    return p.runs[2].text == e["old"]
                if "old_prefix" in e:
                    return p.runs[2].text.startswith(e["old_prefix"])
                return True
            p = find_one(paras, is_target, f"bullet {e['old_lead'][:40]!r}")
            if "new_lead" in e:
                set_run_text(p.runs[1], resolve(e["new_lead"], v, chosen))
            if kind == "bullet":
                set_run_text(p.runs[2], resolve(e["new"], v, chosen))
            else:
                if p.runs[2].text.count(e["find"]) != 1:
                    die(f"bullet_sub: 'find' not found exactly once: {e['find'][:60]!r}")
                set_run_text(p.runs[2], p.runs[2].text.replace(e["find"], resolve(e["new"], v, chosen)))
            log.append(f"{kind}: {e['old_lead'][:40]!r}")
        elif kind == "caption":
            p = find_one(paras, lambda p, o=e["old_prefix"]: p.style.name == "Fig Caption" and p.text.startswith(o),
                         f"caption {e['old_prefix'][:40]!r}")
            new = resolve(e["new"], v, chosen)
            if len(p.runs) == 2 and re.match(r"^FIGURE \d+\.", p.runs[0].text):
                # keep the blue "FIGURE n." label run; the text run follows it
                if not new.startswith(p.runs[0].text):
                    die(f"figure caption replacement must start with its label {p.runs[0].text!r}")
                set_run_text(p.runs[1], new[len(p.runs[0].text):])
            else:
                set_single_run(p, new)
            log.append(f"caption: {e['old_prefix'][:40]!r}")
        elif kind == "cell":
            # locate the table by its header cells, the row by its first cell,
            # the column by header text (or the second column when no col)
            cands = []
            for t in doc.tables:
                h = [c.text for c in t.rows[0].cells]
                if "table_header0" in e and h[0] != e["table_header0"]:
                    continue
                if "table_header1" in e and (len(h) < 2 or h[1] != e["table_header1"]):
                    continue
                cands.append(t)
            if len(cands) != 1:
                die(f"cell edit: expected one table, found {len(cands)} for {e}")
            t = cands[0]
            h = [c.text for c in t.rows[0].cells]
            rows = [r for r in t.rows if r.cells[0].text == e["row"]]
            if len(rows) > 1:
                # several rows share the label (Table III has two "A / C" rows):
                # keep the one that holds the exact old text
                rows = [r for r in rows if any(c.text == e["old"] for c in r.cells)]
            if len(rows) != 1:
                die(f"cell edit: row {e['row']!r} with {e['old']!r} not found once ({len(rows)})")
            row = rows[0]
            if "col" in e:
                ci = h.index(e["col"])
            else:
                ci = next(i for i, c in enumerate(row.cells) if c.text == e["old"])
            if row.cells[ci].text != e["old"]:
                die(f"cell edit: cell text {row.cells[ci].text!r} != {e['old']!r}")
            _set_cell_text(row.cells[ci], resolve(e["new"], v, chosen))
            log.append(f"cell: {e['row'][:30]!r} / {e.get('col', '?')}: {e['old'][:20]!r}")
        elif kind == "insert_bullets_after":
            anchor = find_one(paras, lambda p, o=e["after_prefix"]: p.text.startswith(o),
                              f"insert_bullets_after {e['after_prefix'][:40]!r}")
            if not anchor.text.rstrip().endswith(":"):
                die(f"insert_bullets_after: anchor does not end in a colon: {anchor.text[-40:]!r}")
            nxt = anchor._p.getnext()
            if nxt is not None and nxt.tag == qn("w:p") and "".join(t.text or "" for t in nxt.iter(qn("w:t"))).startswith("• "):
                die("insert_bullets_after: a bullet list already follows the anchor")
            # template: the paragraph formatting and the plain-text run of an
            # existing three-run bullet ("• ", bold lead, text) in the body
            tpls = [p for p in paras if len(p.runs) == 3 and p.runs[0].text == "• " and p.runs[1].bold]
            if not tpls:
                die("insert_bullets_after: no three-run bullet to use as a template")
            tpl = tpls[0]
            from docx.text.paragraph import Paragraph
            prev = anchor._p
            for item in e["items"]:
                new_p = copy.deepcopy(tpl._p)
                prev.addnext(new_p)
                np_ = Paragraph(new_p, anchor._parent)
                runs = np_.runs
                set_run_text(runs[1], "")
                set_run_text(runs[2], resolve(item, v, chosen))
                prev = new_p
            log.append(f"inserted {len(e['items'])} bullets after {e['after_prefix'][:40]!r}")
        elif kind == "delete":
            p = find_one(paras, lambda p, o=e["old_prefix"]: p.text.startswith(o), f"delete {e['old_prefix'][:40]!r}")
            remove_paragraph(p)
            log.append(f"deleted: {e['old_prefix'][:40]!r}")
        else:
            die(f"unknown edit kind {kind}")


# ---------------------------------------------------------------------------
# Table and figure renumbering: VIII -> IX, IX -> X, X -> XI (captions and text)
# ---------------------------------------------------------------------------

def renumber_tables(doc, log):
    """Done BEFORE the edits, on v57 text only, so the edits can use final numbers."""
    shifts = [("X", "XI"), ("IX", "X"), ("VIII", "IX")]       # descending, no collisions
    n = 0
    for p in doc.paragraphs:
        for r in p.runs:
            t = r.text
            for old, new in shifts:
                t2 = re.sub(rf"\bTABLE {old}\b", f"TABLE {new}", t)
                t2 = re.sub(rf"\bTable {old}\b", f"Table {new}", t2)
                if t2 != t:
                    n += 1
                    t = t2
            set_run_text(r, t)
    log.append(f"renumbered Table VIII/IX/X -> IX/X/XI in {n} run(s)")


# ---------------------------------------------------------------------------
# References
# ---------------------------------------------------------------------------

def rework_references(doc, log):
    refs = [p for p in doc.paragraphs if p.style.name == "References"]
    if len(refs) != 56:
        die(f"expected 56 reference paragraphs, found {len(refs)}")
    entries = []     # (old_number or sentinel, paragraph or None, text)
    for p in refs:
        m = re.match(r"\[(\d+)\]\s+", p.text)
        if not m:
            die(f"reference paragraph without a number: {p.text[:50]!r}")
        entries.append([int(m.group(1)), p, p.text])
    for num, text in REF_REPLACE.items():
        hit = [e for e in entries if e[0] == num]
        if len(hit) != 1:
            die(f"reference [{num}] to replace not found once")
        hit[0][2] = text
    for num in REF_DELETE:
        hit = [e for e in entries if e[0] == num]
        if len(hit) != 1:
            die(f"reference [{num}] to delete not found once")
        remove_paragraph(hit[0][1])
        entries.remove(hit[0])
    for after, (sentinel, text) in REF_INSERT_AFTER.items():
        idx = next((i for i, e in enumerate(entries) if e[0] == after), None)
        if idx is None:
            die(f"reference [{after}] to insert after not found")
        anchor = entries[idx][1]
        new_p = copy.deepcopy(anchor._p)
        anchor._p.addnext(new_p)
        from docx.text.paragraph import Paragraph
        np_ = Paragraph(new_p, anchor._parent)
        entries.insert(idx + 1, [sentinel, np_, "[" + sentinel + "]  " + text])
    # new numbering in list order
    mapping = {}
    for i, e in enumerate(entries, start=1):
        mapping[str(e[0])] = i
        body = re.sub(r"^\[[^\]]+\]\s+", "", e[2])
        set_single_run(e[1], f"[{i}]  {body}")
    # in-text citations: every [token] where token is a v57 number or the sentinel
    n = 0
    deleted = {str(d) for d in REF_DELETE}
    def fix(m):
        nonlocal n
        tok = m.group(1)
        if tok in deleted:
            die(f"in-text citation {m.group(0)} refers to a deleted reference; remove the citation first")
        if tok in mapping:
            n += 1
            return f"[{mapping[tok]}]"
        return m.group(0)
    for p in doc.paragraphs:
        if p.style.name == "References":
            continue
        for r in p.runs:
            set_run_text(r, re.sub(r"\[(\d+|LOTTERY)\]", fix, r.text))
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        set_run_text(r, re.sub(r"\[(\d+|LOTTERY)\]", fix, r.text))
    log.append(f"references: {len(entries)} entries, {n} in-text citation token(s) renumbered; "
               f"[45] removed, lottery is [{mapping['LOTTERY']}]")
    return len(entries)


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------

def _set_cell_text(cell, text):
    """Write into the run that carries the cell's formatting. The v57 tables
    hold each cell as an empty, unformatted first run followed by the text
    in an 8-pt run; writing into the first run would render at the body
    size (seen on the first dry build)."""
    ps = cell.paragraphs
    runs = ps[0].runs
    if not runs:
        ps[0].add_run(text)
    else:
        target = next((r for r in runs if r._r.find(qn("w:rPr")) is not None), runs[-1])
        for r in runs:
            if r is not target:
                set_run_text(r, "")
        set_run_text(target, text)
    for extra in ps[1:]:
        extra._p.getparent().remove(extra._p)


def _set_grid(tbl, widths):
    grid = tbl.find(qn("w:tblGrid"))
    for g in list(grid):
        grid.remove(g)
    from docx.oxml import OxmlElement
    for w in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(w)); grid.append(gc)
    for tr in tbl.findall(qn("w:tr")):
        for tc, w in zip(tr.findall(qn("w:tc")), widths):
            tcpr = tc.find(qn("w:tcPr"))
            tcw = tcpr.find(qn("w:tcW")) if tcpr is not None else None
            if tcw is not None:
                tcw.set(qn("w:w"), str(w)); tcw.set(qn("w:type"), "dxa")


def edit_table_vii(doc, v, chosen, log):
    t = doc.tables[6]
    if [c.text for c in t.rows[0].cells] != ["", "A", "+ concurrency", "+ retry", "+ log table", "+ ledger"]:
        die("Table VII header is not the v57 one")
    # insert the O column after A by cloning the A cell in every row
    for row in t.rows:
        tcs = row._tr.findall(qn("w:tc"))
        clone = copy.deepcopy(tcs[1])
        tcs[1].addnext(clone)
    t = doc.tables[6]
    for row in t.rows:
        if len(row.cells) != 7:
            die("Table VII column insert failed")
    total = 10080
    label_w = 2419
    other = (total - label_w) // 6
    _set_grid(t._tbl, [label_w] + [other] * 6)
    for i, h in enumerate(TABLE_VII["header"]):
        _set_cell_text(t.rows[0].cells[i], h)
    for row in t.rows[1:]:
        key = row.cells[0].text.strip()
        spec = next((vals for k, vals in TABLE_VII["rows"].items() if key.startswith(k)), None)
        if spec is None:
            die(f"Table VII row label not in TABLE_VII spec: {key!r}")
        for i, val in enumerate(spec, start=1):
            _set_cell_text(row.cells[i], resolve(val, v, chosen))
    log.append("Table VII: O column inserted, 7 columns, all cells filled")
    return t


def insert_table_viii(doc, table_vii, v, chosen, log):
    """Clone Table VII (now 7 columns) and its two caption paragraphs into a
    one-column section of its own, placed right after Figure 6's section:
    [label, caption, table, sectPr(cols=1)], which is how the template holds
    every full-width table. Appending it inside Table VII's section put it
    under Figure 6's floating image in LibreOffice's rendering."""
    tbl = table_vii._tbl
    body = tbl.getparent()
    kids = list(body)
    i = kids.index(tbl)
    label_p, cap_p = kids[i - 2], kids[i - 1]
    if "".join(x.text or "" for x in label_p.iter(qn("w:t"))).strip() != "TABLE VII":
        die("could not find the 'TABLE VII' label paragraph above the table")
    # Figure 6's section: caption paragraph, then the paragraph carrying its sectPr
    fig_cap = next((k for k in kids if k.tag == qn("w:p")
                    and "".join(x.text or "" for x in k.iter(qn("w:t"))).startswith("FIGURE 6.")), None)
    if fig_cap is None:
        die("Figure 6 caption not found")
    j = kids.index(fig_cap)
    sect_p = None
    for k in kids[j + 1:j + 4]:
        ppr = k.find(qn("w:pPr")) if k.tag == qn("w:p") else None
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            sect_p = k
            break
    if sect_p is None:
        die("no section-ending paragraph found after the Figure 6 caption")
    cols = sect_p.find(qn("w:pPr")).find(qn("w:sectPr")).find(qn("w:cols"))
    if cols is not None and cols.get(qn("w:num")) not in (None, "1"):
        die("the section after Figure 6 is not one-column")
    new_label, new_cap, new_tbl, new_sect = (copy.deepcopy(label_p), copy.deepcopy(cap_p),
                                             copy.deepcopy(tbl), copy.deepcopy(sect_p))
    trs = new_tbl.findall(qn("w:tr"))
    for tr in trs[4:]:
        new_tbl.remove(tr)
    sect_p.addnext(new_label); new_label.addnext(new_cap); new_cap.addnext(new_tbl); new_tbl.addnext(new_sect)
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    lp, cp = Paragraph(new_label, table_vii._parent), Paragraph(new_cap, table_vii._parent)
    set_single_run(lp, TABLE_VIII["label"])
    set_single_run(cp, resolve(TABLE_VIII["caption"], v, chosen))
    nt = Table(new_tbl, table_vii._parent)
    for i, h in enumerate(TABLE_VIII["header"]):
        _set_cell_text(nt.rows[0].cells[i], h)
    for r, vals in zip(nt.rows[1:], TABLE_VIII["rows"]):
        for i, val in enumerate(vals):
            _set_cell_text(r.cells[i], resolve(val, v, chosen))
    log.append("Table VIII (p sweep) inserted in its own one-column section after Figure 6")


def rebuild_table_xi(doc, reports, v, log):
    """The per-seed table (v57's Table X, now XI): 50 rows x 8 columns from
    the per-trial lists in the p = 0.30 report."""
    cands = [t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Seed"]
    if len(cands) != 1:
        die(f"expected exactly one per-seed table (header 'Seed'), found {len(cands)}")
    t = cands[0]
    tbl_el = t._tbl
    blk = reports["p30"][1]["configs"]
    cols = [("A", "A"), ("O", "O"), ("+ concurrency", "C no retry"), ("+ retry", "C retry"),
            ("+ log table", "C log"), ("+ ledger", "C ledger")]
    per = {label: {t_["seed"]: t_ for t_ in blk[label]["trials"]} for _, label in cols}
    seeds = sorted(set().union(*[set(d) for d in per.values()]))
    # make the header 8 columns by cloning the last header cell as needed
    tr0 = t.rows[0]._tr
    while len(tr0.findall(qn("w:tc"))) < 8:
        tr0.append(copy.deepcopy(tr0.findall(qn("w:tc"))[-1]))
    while len(tr0.findall(qn("w:tc"))) > 8:
        tr0.remove(tr0.findall(qn("w:tc"))[-1])
    # one template data row, then rebuild
    template = copy.deepcopy(t.rows[1]._tr)
    while len(template.findall(qn("w:tc"))) < 8:
        template.append(copy.deepcopy(template.findall(qn("w:tc"))[-1]))
    while len(template.findall(qn("w:tc"))) > 8:
        template.remove(template.findall(qn("w:tc"))[-1])
    for tr in tbl_el.findall(qn("w:tr"))[1:]:
        tbl_el.remove(tr)
    for s in seeds:
        tbl_el.append(copy.deepcopy(template))
    from docx.table import Table
    t = Table(tbl_el, t._parent)
    heads = ["Seed", "Layers with an injected failure"] + [c for c, _ in cols]
    for i, h in enumerate(heads):
        _set_cell_text(t.rows[0].cells[i], h)
    # The injected layers are a pure function of (scale, seed) under the
    # paired schedule, so they are computed, not read off Scenario A's
    # outcome -- an excluded A trial would otherwise print "none".
    from scenarios.common import faults
    from experiments.run_experiment import fault_seed
    p30 = float(reports["p30"][1]["configs"]["A"]["fault_p"][0])
    for row, s in zip(t.rows[1:], seeds):
        fs = fault_seed(10000, s)
        inj = [l for l in ("relational", "nosql", "vector") if faults.paired_fault_u(fs, l, 0) < p30]
        a = per["A"].get(s)
        if a is not None and sorted(a["leaked_layers"]) != sorted(inj):
            die(f"Table XI: Scenario A seed {s} leaked {a['leaked_layers']} but the schedule injects {inj}")
        injected = ", ".join(inj) if inj else "none"
        vals = [str(s), injected]
        for _, label in cols:
            tr_ = per[label].get(s)
            if tr_ is None:
                vals.append("excluded")
            elif tr_["leaked_layers"]:
                vals.append("unbounded")
            else:
                vals.append(f"{tr_['drift_window_s']:.3f}")
        for i, val in enumerate(vals):
            _set_cell_text(row.cells[i], val)
    total = 10080
    w0, w1 = 700, 2180
    other = (total - w0 - w1) // 6
    _set_grid(t._tbl, [w0, w1] + [other] * 6)
    log.append(f"Table XI rebuilt: {len(seeds)} seeds x 8 columns")


# ---------------------------------------------------------------------------
# Figure 6
# ---------------------------------------------------------------------------

def _png_size(blob):
    if blob[:8] != b"\x89PNG\r\n\x1a\n":
        die("figure is not a PNG")
    w, h = struct.unpack(">II", blob[16:24])
    return w, h


def replace_figure6(doc, png_path, log):
    paras = doc.paragraphs
    cap = find_one(paras, lambda p: p.text.startswith("FIGURE 6."), "Figure 6 caption")
    idx = paras.index(cap)
    img_p = paras[idx - 1]
    blips = list(img_p._p.iter(qn("a:blip")))
    if len(blips) != 1:
        die("paragraph above the Figure 6 caption does not hold exactly one image")
    rid = blips[0].get(qn("r:embed"))
    part = doc.part.related_parts[rid]
    blob = Path(png_path).read_bytes()
    w, h = _png_size(blob)
    part._blob = blob
    for ext in list(img_p._p.iter(qn("wp:extent"))) + list(img_p._p.iter(qn("a:ext"))):
        cx = int(ext.get("cx"))
        ext.set("cy", str(int(cx * h / w)))
    log.append(f"Figure 6 image replaced ({w}x{h} px, {len(blob)} bytes), height rescaled to its aspect")


# ---------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------

def check(doc, n_refs, log):
    problems = []
    texts = [p.text for p in doc.paragraphs]
    joined = "\n".join(texts)
    markers = re.findall(r"\{\{[^}]*\}\}|\[\[[^\]]*\]\]", joined)
    if markers:
        problems.append(f"unresolved placeholder or marker left in the text: {markers[:5]}")
    for phrase in ("comparable in rate, not paired", "failure positions", "separate generators"):
        hits = [t[:70] for t in texts if phrase in t]
        allowed = 1 if phrase == "separate generators" else 0
        if len(hits) > allowed:
            problems.append(f"stale phrase {phrase!r} in {len(hits)} paragraph(s): {hits}")
    ab = next(p for p in doc.paragraphs if p.style.name == "Abstract")
    words = len(ab.runs[1].text.split())
    log.append(f"abstract: {words} words")
    if words > 250:
        problems.append(f"abstract is {words} words (> 250)")
    # citations: ascending first appearance, all present
    first = {}
    for i, t in enumerate(texts):
        if doc.paragraphs[i].style.name == "References":
            continue
        for m in re.finditer(r"\[(\d+)\]", t):
            first.setdefault(int(m.group(1)), i)
    order = [k for k, _ in sorted(first.items(), key=lambda kv: (kv[1], kv[0]))]
    if order != sorted(order):
        problems.append(f"citations not in ascending first-use order: {order[:12]}...")
    missing = sorted(set(range(1, n_refs + 1)) - set(first))
    if missing:
        problems.append(f"references never cited: {missing}")
    if max(first) > n_refs:
        problems.append(f"citation beyond the reference list: {max(first)} > {n_refs}")
    # captions sequence
    tabs = [t for t in texts if re.fullmatch(r"TABLE [IVX]+", t.strip())]
    want = [f"TABLE {ROMAN[i]}" for i in range(len(tabs))]
    if tabs != want:
        problems.append(f"table labels out of sequence: {tabs}")
    figs = [re.match(r"FIGURE (\d+)\.", t).group(1) for t in texts if re.match(r"FIGURE \d+\.", t)]
    if figs != [str(i) for i in range(1, len(figs) + 1)]:
        problems.append(f"figure captions out of sequence: {figs}")
    log.append(f"{len(tabs)} tables, {len(figs)} figures, {n_refs} references, citations ascending")
    return problems


def all_text(doc):
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            out.append(" | ".join(c.text for c in r.cells))
    return "\n".join(out)


# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--v57", required=True)
    ap.add_argument("--out")
    ap.add_argument("--p30"); ap.add_argument("--p10"); ap.add_argument("--p05")
    ap.add_argument("--fig6", help="PNG of the six-bar Figure 6 from make_paper_figures.py")
    ap.add_argument("--roundtrip", action="store_true", help="open and save with no edits; verify text identical")
    ap.add_argument("--allow-markers", action="store_true",
                    help="dry builds only: write the file even if a [[marker]] is left (never for a submission)")
    args = ap.parse_args()

    doc = Document(args.v57)
    if args.roundtrip:
        before = all_text(doc)
        out = args.out or "/tmp/roundtrip.docx"
        doc.save(out)
        after = all_text(Document(out))
        print("roundtrip:", "IDENTICAL" if before == after else "DIFFERS")
        sys.exit(0 if before == after else 1)

    if not (args.p30 and args.p10 and args.p05 and args.out):
        die("--p30, --p10, --p05 and --out are required for a build")
    v, reports = load_values(args.p30, args.p10, args.p05)
    log, chosen = [], []

    renumber_tables(doc, log)               # v57 text only, before edits
    apply_edits(doc, v, chosen, log)
    tvii = edit_table_vii(doc, v, chosen, log)
    insert_table_viii(doc, tvii, v, chosen, log)
    rebuild_table_xi(doc, reports, v, log)
    if args.fig6:
        replace_figure6(doc, args.fig6, log)
    else:
        log.append("Figure 6 NOT replaced (no --fig6)")
    n_refs = rework_references(doc, log)

    for c in chosen:
        print("  ", c)
    for l in log:
        print("  ", l)
    problems = check(doc, n_refs, log)
    if args.allow_markers:
        problems = [p for p in problems if not p.startswith("unresolved placeholder")]
        print("   (--allow-markers: marker check waived; this file is NOT for submission)")
    if problems:
        for p in problems:
            print("  PROBLEM:", p)
        die("checks failed; nothing written")
    doc.save(args.out)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
