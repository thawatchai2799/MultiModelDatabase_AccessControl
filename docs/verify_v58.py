#!/usr/bin/env python3
"""Independent check of the v58 manuscript against the raw campaign files.

    python3 docs/verify_v58.py --docx MultiModelDatabase_AccessControl_IEEE_v58.docx \\
        --results /path/to/paired/results

Everything here is recomputed from the per-layer jsonl rows -- NOT from the
ablation_table.py JSON report and NOT from fill_revision.py -- so a mistake
shared by the report and the paper would still be caught. The checks are:

  1. Table VII  (p = 0.30): k of n, Wilson 95% interval, median window of the
     closed trials, median caller-observed revoke latency, confirmed
     containment k_closed of n, leak attribution, per column.
  2. Table VIII (p sweep): k of n, Wilson, pairs a/b and exact McNemar p.
  3. Table XI   (per seed): injected layers, every cell (window / unbounded /
     excluded) for the 50 seeds x 6 configurations.
  4. Prose: every "k of n" and "[lo, hi]" token in the body text that refers
     to the campaign must be one the data can produce; anything else is
     printed with its context for a human to look at.

Exit status is non-zero on any mismatch.
"""
import argparse
import json
import math
import re
import statistics
import sys
from collections import defaultdict
from pathlib import Path

from docx import Document

LAYERS = ("relational", "nosql", "vector")
TIMEOUT = 30.0
COLS = [("A", "none-noretry", "a"), ("O", "none-noretry", "o"), ("+ concurrency", "none-noretry", "c"),
        ("+ retry", "none", "c"), ("+ log table", "log", "c"), ("+ ledger", "ledger", "c")]
PS = ["0.05", "0.10", "0.30"]

bad = []


def fail(msg):
    bad.append(msg)
    print("MISMATCH:", msg)


def wilson(k, n, z=1.959964):
    if n == 0:
        return (0.0, 0.0)
    p = k / n
    den = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / den
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return (max(0.0, centre - half), min(1.0, centre + half))


def mcnemar_exact(b, c):
    n = b + c
    if n == 0:
        return 1.0
    k = min(b, c)
    tail = sum(math.comb(n, i) for i in range(0, k + 1)) / 2 ** n
    return min(1.0, 2 * tail)


def load_trials(path, scenario):
    """One dict per trial, recomputed from the layer rows."""
    by = defaultdict(dict)
    for line in Path(path).read_text().splitlines():
        if not line.strip():
            continue
        r = json.loads(line)
        if r.get("scenario") != scenario or r.get("layer") not in LAYERS:
            continue
        key = (r["seed"], r["resource_id"], r["principal_id"], r["t_issued"])
        by[key][r["layer"]] = r
    trials = {}
    for key, layers in by.items():
        seed = key[0]
        if seed in trials:
            fail(f"{path}: seed {seed} appears twice for scenario {scenario}")
        if set(layers) != set(LAYERS):
            trials[seed] = {"excluded": "incomplete"}
            continue
        ex = [l["extra"] for l in layers.values()]
        if any(e.get("retrievable_before_revoke") is False for e in ex):
            trials[seed] = {"excluded": "invalid_setup"}
            continue
        if any(e.get("poller_error") for e in ex):
            trials[seed] = {"excluded": "poller_error"}
            continue
        leaked = sorted(l for l, rec in layers.items() if rec["confirmed_contained"] is not True)
        wins = [l["leak_window_s"] for l in layers.values()]
        e0 = layers[LAYERS[0]]["extra"]
        t = {
            "excluded": None,
            "leaked": bool(leaked), "leaked_layers": leaked,
            "window": None if (leaked or any(w is None for w in wins)) else max(wins),
            "latency": e0.get("revoke_latency_s"),
            "injected": None,
            "status": None,
        }
        if scenario == "a":
            # the application records which first attempts it lost
            t["injected"] = [l for l in LAYERS if l in (e0.get("injected_failed_layers") or [])]
        elif scenario == "c":
            # the bridge reports attempts per layer: a first-attempt failure
            # shows as attempts > 1 (retry on) or ok == False (retry off)
            prop = {x["layer"]: x for x in (e0.get("bridge_response") or {}).get("propagation", [])}
            t["injected"] = [l for l in LAYERS if l in prop and (prop[l].get("attempts", 1) > 1 or prop[l].get("ok") is False)]
        elif scenario == "o":
            # relational is written inline with the outbox row; nosql and
            # vector by the relay worker
            inl = (e0.get("outbox_revoke") or {}).get("inline_attempts", 1)
            wl = (e0.get("outbox_worker") or {}).get("layers") or {}
            t["injected"] = [l for l in LAYERS if (l == "relational" and inl > 1)
                             or (l != "relational" and wl.get(l, {}).get("attempts", 1) > 1)]
        if scenario == "c":
            for l in reversed(LAYERS):
                st = layers[l]["extra"].get("bridge_status")
                if isinstance(st, dict) and "contained" in st:
                    t["status"] = st
                    break
        elif scenario == "o":
            t["status"] = e0.get("outbox_status")
        trials[seed] = t
    return trials


def summarise(trials):
    valid = {s: t for s, t in trials.items() if t["excluded"] is None}
    n = len(valid)
    k = sum(t["leaked"] for t in valid.values())
    closed = [t["window"] for t in valid.values() if not t["leaked"]]
    lat = [t["latency"] for t in valid.values() if t["latency"] is not None]
    confirmed = sum(1 for t in valid.values() if t["status"] and t["status"].get("contained") is True)
    false_cont = sum(1 for t in valid.values() if t["status"] and t["status"].get("contained") is True and t["leaked"])
    attributed = sum(1 for t in valid.values() if t["leaked"] and t["status"] and t["status"].get("contained") is False
                     and sorted(t["status"].get("missingLayers") or []) == t["leaked_layers"])
    return {
        "n": n, "k": k, "ci": wilson(k, n),
        "win": statistics.median(closed) if closed else None,
        "lat": statistics.median(lat) if lat else None,
        "confirmed": confirmed, "false_cont": false_cont, "attributed": attributed,
        "leaked_seeds": sorted(s for s, t in valid.items() if t["leaked"]),
        "has_status": any(t["status"] for t in valid.values()),
    }


def table_rows(t):
    out = []
    for r in t.rows:
        cells, prev = [], None
        for c in r.cells:
            if c._tc is not prev:
                cells.append(c.text.strip())
            prev = c._tc
        out.append(cells)
    return out


def find_table(doc, header0, ncols=None, header1=None):
    hits = [t for t in doc.tables if t.rows[0].cells[0].text.strip() == header0
            and (ncols is None or len(table_rows(t)[0]) == ncols)
            and (header1 is None or table_rows(t)[0][1] == header1)]
    if len(hits) != 1:
        raise SystemExit(f"table with header {header0!r} not found once ({len(hits)})")
    return hits[0]


def eq(label, got, want):
    if got != want:
        fail(f"{label}: paper says {got!r}, data says {want!r}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", required=True)
    ap.add_argument("--results", required=True, help="directory with paired_p*_*.jsonl")
    a = ap.parse_args()
    res = Path(a.results)
    doc = Document(a.docx)

    # ---- recompute everything -------------------------------------------
    data = {}   # (p, col) -> (trials, summary)
    for p in PS:
        for col, stem, scen in COLS:
            f = res / f"paired_p{p}_{stem}.jsonl"
            tr = load_trials(f, scen)
            data[(p, col)] = (tr, summarise(tr))
            print(f"p={p} {col:14s} n={data[(p,col)][1]['n']} k={data[(p,col)][1]['k']}")

    # ---- 1. Table VII ----------------------------------------------------
    t7 = {r[0]: r[1:] for r in table_rows(find_table(doc, "", 7)) if r and r[0]}
    hdr = table_rows(find_table(doc, "", 7))[0][1:]
    eq("Table VII header", hdr, [c for c, _, _ in COLS])
    for i, (col, _, scen) in enumerate(COLS):
        tr, s = data[("0.30", col)]
        lo, hi = s["ci"]
        eq(f"Table VII leak {col}", t7["Trials leaking without bound"][i], f"{s['k']} of {s['n']} [{lo:.2f}, {hi:.2f}]")
        cell = t7["Drift window (median of the trials that closed)"][i].replace(" †", "")
        eq(f"Table VII window {col}", cell, f"{s['win']:.3f} s")
        eq(f"Table VII latency {col}", t7["Revoke latency (median)"][i], f"{s['lat']:.3f} s")
        conf_row = [k for k in t7 if k.startswith("Trials the record could confirm")][0]
        want = "n/a" if scen == "a" else f"{s['confirmed']} of {s['n']}"
        eq(f"Table VII confirm {col}", t7[conf_row][i], want)
        if s["false_cont"]:
            fail(f"{col}: {s['false_cont']} false containment(s) in data")
        att = t7["Leak detected and attributed"][i]
        if s["has_status"] and s["k"] > 0:
            eq(f"Table VII attributed {col}", att, f"{s['attributed']} of {s['k']}")
        elif s["has_status"] and s["k"] == 0:
            eq(f"Table VII attributed {col}", att, "no leak to attribute")
        else:
            eq(f"Table VII attributed {col}", att, "never")
    # A's dagger: closed-trial count
    sA = data[("0.30", "A")][1]
    print(f"A closed trials = {sA['n'] - sA['k']}, leaked = {sA['k']}")

    # ---- 2. Table VIII ---------------------------------------------------
    t8 = table_rows(find_table(doc, "p", 7))
    eq("Table VIII header", t8[0][1:], [c for c, _, _ in COLS])
    for row in t8[1:]:
        p = row[0]
        trA, sA = data[(p, "A")]
        for i, (col, _, _) in enumerate(COLS):
            tr, s = data[(p, col)]
            lo, hi = s["ci"]
            want = f"{s['k']} of {s['n']} [{lo:.2f}, {hi:.2f}]"
            if col != "A":
                common = [sd for sd in trA if trA[sd]["excluded"] is None and sd in tr and tr[sd]["excluded"] is None]
                b = sum(1 for sd in common if trA[sd]["leaked"] and not tr[sd]["leaked"])
                c = sum(1 for sd in common if tr[sd]["leaked"] and not trA[sd]["leaked"])
                pv = mcnemar_exact(b, c)
                ptxt = "p < 0.001" if pv < 0.001 else f"p = {pv:.3f}"
                want += f"; pairs {b}/{c}, {ptxt}"
            eq(f"Table VIII p={p} {col}", row[1 + i], want)

    # ---- 3. Table XI -----------------------------------------------------
    t11 = table_rows(find_table(doc, "Seed"))
    eq("Table XI header", t11[0], ["Seed", "Layers with an injected failure"] + [c for c, _, _ in COLS])
    eq("Table XI rows", len(t11) - 1, 50)
    for row in t11[1:]:
        seed = int(row[0])
        inj = None
        for col, _, _ in COLS:
            t = data[("0.30", col)][0].get(seed)
            if t and t["excluded"] is None:
                if inj is None:
                    inj = t["injected"]
                elif t["injected"] != inj:
                    fail(f"seed {seed}: injected layers differ between columns ({inj} vs {t['injected']})")
        eq(f"Table XI seed {seed} injected", row[1], ", ".join(inj) if inj else "none")
        for i, (col, _, _) in enumerate(COLS):
            t = data[("0.30", col)][0].get(seed)
            if t is None:
                want = "missing"
            elif t["excluded"]:
                want = "excluded"
            elif t["leaked"]:
                want = "unbounded"
            else:
                want = f"{t['window']:.3f}"
            eq(f"Table XI seed {seed} {col}", row[2 + i], want)
        # A must leak exactly when something was injected
        tA = data[("0.30", "A")][0].get(seed)
        if tA and tA["excluded"] is None and tA["leaked"] != bool(inj):
            fail(f"seed {seed}: A leaked={tA['leaked']} but injected={inj}")

    # ---- 3a. Tables V and VI: the 10K faulty cells and the derived extras --
    sA, sCr, sClog, sCled = (data[("0.30", c)][1] for c in ("A", "+ retry", "+ log table", "+ ledger"))
    tA, tCled = data[("0.30", "A")][0], data[("0.30", "+ ledger")][0]
    n_inj_seeds = sum(1 for t in tA.values() if t["excluded"] is None and t["injected"]) \
        + sum(1 for sd, t in tA.items() if t["excluded"] and data[("0.30", "+ retry")][0][sd]["injected"])
    t5 = {r[0]: r[1:] for r in table_rows(find_table(doc, "", header1="1K (5 seeds)"))}
    lo, hi = sA["ci"]
    eq("Table V hit", t5["Trials in which at least one layer was hit"][1],
       f"{n_inj_seeds} of {len(tA)} ({sA['k']} of {sA['n']} valid) §")
    eq("Table V leaked", t5["Trials leaking past the 30 s limit"][1],
       f"{sA['k']} of {sA['n']} = {sA['k'] / sA['n']:.2f}, Wilson 95% CI [{lo:.2f}, {hi:.2f}]")
    eq("Table V hit->leaked", t5["Of the trials that were hit, how many leaked without bound"][1], f"{sA['k']} of {sA['k']}")
    eq("Table V window", t5["Median window of the trials that closed ‖"][1], f"{sA['win'] * 1000:.0f} ms")
    t6 = {r[0]: r[1:] for r in table_rows(find_table(doc, "", header1="A healthy"))}
    drow = [k for k in t6 if k.startswith("Drift window observed at 10 ms polling")][0]
    eq("Table VI A faulty window 10K", t6[drow][2].split(" / ")[-1], f"{sA['win']:.3f} s")
    eq("Table VI C faulty window 10K", t6[drow][3].split(" / ")[-1], f"{sCled['win']:.2f} s")
    eq("Table VI C faulty bound 10K", t6["Leak Window, system bound (10K) †"][3], f"{sCled['win']:.2f} s")
    lo2, hi2 = sCled["ci"]
    eq("Table VI A unbounded", t6["Unbounded leaks at 10K"][2], f"{sA['k']} of {sA['n']} [{lo:.2f}, {hi:.2f}]")
    eq("Table VI C unbounded", t6["Unbounded leaks at 10K"][3], f"{sCled['k']} of {sCled['n']} [{lo2:.2f}, {hi2:.2f}]")
    eq("Table VI C false containment", t6["False containment"][3], f"{sCled['false_cont']} of {sCled['n']}")
    body_all = "\n".join(p.text for p in doc.paragraphs)
    extra_led = f"{sCled['lat'] - sClog['lat']:.1f} s"
    extra_log = f"{sClog['lat'] - sCr['lat']:.2f} s"
    for tok in (extra_led, extra_log):
        if tok not in body_all:
            fail(f"derived latency difference {tok!r} not found in the text")
    print(f"ledger extra {extra_led}, log extra {extra_log}: present")

    # ---- 3b. Exclusion label in the Table XI caption -----------------------
    cap = next((p.text for p in doc.paragraphs if p.text.startswith("Every trial of the fifty-seed faulty grid")), "")
    for col, stem, scen in COLS:
        for line in (res / f"paired_p0.30_{stem}.jsonl").read_text().splitlines():
            r = json.loads(line) if line.strip() else {}
            if r.get("scenario") == scen and (r.get("extra") or {}).get("poller_error"):
                msg = r["extra"]["poller_error"]
                want = f"{col} seed {r['seed']} at p = 0.30 (clock anomaly)" if "clock anomaly" in msg else None
                if want is None or want not in cap:
                    fail(f"Table XI caption {cap[-80:]!r} does not describe the excluded trial: {col} seed {r['seed']}: {msg!r}")

    # ---- 4. Prose tokens -------------------------------------------------
    legit = set()
    for (p, col), (tr, s) in data.items():
        lo, hi = s["ci"]
        legit.add(f"{s['k']} of {s['n']}")
        legit.add(f"{s['n'] - s['k']} of {s['n']}")
        legit.add(f"[{lo:.2f}, {hi:.2f}]")
        lo2, hi2 = wilson(s["n"] - s["k"], s["n"])
        legit.add(f"[{lo2:.2f}, {hi2:.2f}]")
        if s["k"]:
            legit.add(f"{s['attributed']} of {s['k']}")
    def newcombe(k1, n1, k2, n2):
        p1, p2 = k1 / n1, k2 / n2
        l1, u1 = wilson(k1, n1)
        l2, u2 = wilson(k2, n2)
        d = p1 - p2
        return d - math.sqrt((p1 - l1) ** 2 + (u2 - p2) ** 2), d + math.sqrt((u1 - p1) ** 2 + (p2 - l2) ** 2)
    for p in PS:
        ss = {c: data[(p, c)][1] for c, _, _ in COLS}
        for a_ in ss:
            for b_ in ss:
                if a_ != b_:
                    lo, hi = newcombe(ss[a_]["k"], ss[a_]["n"], ss[b_]["k"], ss[b_]["n"])
                    legit.add(f"[{lo:+.2f}, {hi:+.2f}]")
    body = "\n".join(p.text for p in doc.paragraphs)
    refs_at = body.find("REFERENCES")
    body = body[:refs_at] if refs_at > 0 else body
    print("\nProse tokens to eyeball (not produced by the 50-seed campaign):")
    for m in re.finditer(r"\b\d+ of \d+\b|\[[+−-]?\d\.\d\d, [+−-]?\d\.\d\d\]", body):
        tok = m.group(0)
        if tok not in legit:
            ctx = body[max(0, m.start() - 70):m.end() + 40].replace("\n", " ")
            print(f"   {tok!r:22s} ... {ctx}")

    print("\nRESULT:", "OK, 0 mismatches" if not bad else f"{len(bad)} mismatch(es)")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
